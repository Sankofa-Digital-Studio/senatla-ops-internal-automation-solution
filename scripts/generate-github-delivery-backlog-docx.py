from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "output" / "doc" / "Senatla-Ops-GitHub-Delivery-Backlog.docx"
NAVY = "17365D"
PALE_BLUE = "EAF1F8"
LIGHT_GRAY = "F5F7FA"
BORDER = "D9D9D9"


TASKS = [
    (86, "P1", "10", "Persist the complete site day record", "3 to 5 days", "M1 Operational Foundations", [
        "Persist work plan, location, trips, safety, attendance, signature reference, and delivery result in one versioned payload.",
        "Reuse the existing idempotency key and queue so retries cannot duplicate attendance or trip records.",
        "Reject incomplete context and sites outside the actor's permitted scope.",
        "Prove online, retryable, and accepted outcomes on desktop and at 390 px.",
    ]),
    (87, "P0", "4", "Make dev the integration default and protect main releases", "0.5 to 1 day", "M0 Governance and Access", [
        "Set dev as the default branch and require pull requests for protected branches.",
        "Require the agreed application, database, responsive, and native checks.",
        "Require CODEOWNER approval for main and prevent force pushes and deletion.",
        "Prove the rules with a test release pull request and align contributor documentation.",
    ]),
    (88, "P0", "7", "Verify hosted runtime configuration and preview smoke tests", "1 to 2 days", "M0 Governance and Access", [
        "Confirm Vercel team, project, and branch ownership without recording secrets.",
        "Fail hosted builds when required public Supabase configuration is missing.",
        "Prove login, one protected route, and unauthenticated redirection against an exact commit.",
        "Record the current deployment, environment, responsible owner, and rollback target.",
    ]),
    (89, "P1", "11", "Enforce ownership SLA and immutable issue history", "3 to 4 days", "M1 Operational Foundations", [
        "Replace mutable audit JSON with append-only issue events.",
        "Enforce valid statuses, transitions, owners, due dates, and organization and site access.",
        "Prevent clients from changing or deleting history.",
        "Cover creation, escalation, resolution, and prohibited changes in UI and pgTAP tests.",
    ]),
    (90, "P1", "13", "Convert failed inspections into defects and safety blocks", "5 to 8 days", "M2 Controlled Pilot", [
        "Add versioned inspection templates, runs, answers, defects, and safety blocks.",
        "Create a defect and work order atomically when a critical item fails.",
        "Prevent blocked assets from assignment or return to service.",
        "Prove mobile capture, retry safety, and authorization with pgTAP and 390 px tests.",
    ]),
    (91, "P1", "12", "Stage asset CSV imports before writes", "4 to 6 days", "M2 Controlled Pilot", [
        "Provide a versioned template and dry run with no writes.",
        "Report missing identifiers, duplicates, stored conflicts, invalid links, and invalid custodians or sites.",
        "Apply only an explicitly approved clean import in one transaction.",
        "Prove case-insensitive identity, cycle prevention, audit history, and idempotent re-import.",
    ]),
    (92, "P1", "14", "Replace unsupported executive claims", "3 to 5 days", "M2 Controlled Pilot", [
        "Define source, period, unit, numerator, denominator, and calculation version for every metric.",
        "Remove hard-coded success claims and distinguish actual, target, provisional, stale, and missing values.",
        "Provide drill-down to contributing source records.",
        "Reconcile empty, partial, duplicate, provisional, and complete fixtures in integer minor units.",
    ]),
    (93, "P1", "13", "Enforce independent return to service approval", "2 to 4 days", "M2 Controlled Pilot", [
        "Move blocker checks and activation into a restricted database function.",
        "Reject self-approval, wrong-role access, stale requests, and unresolved safety or maintenance blockers.",
        "Make approval and activation one atomic transaction.",
        "Record the decision and asset state transition in immutable audit events.",
    ]),
    (94, "P1", "15", "Move operational evidence into private storage", "4 to 6 days", "M1 Operational Foundations", [
        "Replace base64 and typed paths with private object references and controlled metadata.",
        "Authorize short-lived access through role and site checks.",
        "Retain hash, content type, size, actor, source record, and retention status.",
        "Prove cross-site denial, retry idempotency, and safe orphan cleanup without leaking paths.",
    ]),
    (95, "P2", "16", "Generate and verify approved documents with one command", "3 to 5 days", "M1 Operational Foundations", [
        "Add an artifact manifest and one documented generation command.",
        "Fail on missing output, render failure, broken link, version mismatch, or prohibited sensitive patterns.",
        "Produce no source-controlled difference on a second run.",
        "Separate approved, internal, and superseded artifacts and retain a compact CI failure report.",
    ]),
    (96, "P0", "17", "Add release journeys for every operational role", "4 to 7 days", "M3 UAT and Release", [
        "Cover Site Manager, Office Administrator, Director, and read-only critical journeys on desktop and at 390 px.",
        "Verify direct-route and control denial for unauthorized roles.",
        "Retain screenshot, trace, route, role, viewport, and commit SHA on failure.",
        "Use approved fixtures without client personal data and link evidence from the release checklist.",
    ]),
    (97, "P1", "20", "Publish a reusable application quality workflow", "2 to 4 days", "M4 Reuse and Scale", [
        "Extract pinned Node, package, Supabase, browser, permission, timeout, and artifact controls behind workflow inputs.",
        "Keep project-specific commands configurable and cache by compatible lockfile.",
        "Prove Senatla Ops can consume the workflow without weakening its current gate.",
        "Validate a second consumer example and publish an adoption checklist.",
    ]),
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = BORDER) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), color)


def set_cell_margin(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + name))
        if node is None:
            node = OxmlElement("w:" + name)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, separate, end])


def add_bullet(doc: Document, text: str, keep_with_next: bool = False) -> None:
    p = doc.add_paragraph(text, style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = keep_with_next


def remove_paragraph_borders(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def add_heading(doc: Document, text: str, level: int) -> None:
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True


def build_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.76)
    section.right_margin = Inches(0.76)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12

    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(29)
    styles["Title"].font.bold = True
    styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    title_style_ppr = styles["Title"].element.get_or_add_pPr()
    title_style_border = title_style_ppr.find(qn("w:pBdr"))
    if title_style_border is not None:
        title_style_ppr.remove(title_style_border)
    styles["Subtitle"].font.name = "Aptos"
    styles["Subtitle"].font.size = Pt(13)
    styles["Subtitle"].font.color.rgb = RGBColor(60, 60, 60)
    for name, size in (("Heading 1", 18), ("Heading 2", 13.5), ("Heading 3", 11.5)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    doc.core_properties.title = "Senatla Ops GitHub Delivery Backlog"
    doc.core_properties.subject = "Repository audit, project board structure, and implementation-ready delivery tasks"
    doc.core_properties.author = "Sankofa Digital Studio"
    doc.core_properties.keywords = "Senatla Ops, GitHub, backlog, project board, delivery, UAT"

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_p.add_run("Senatla Ops  |  GitHub Delivery Backlog  |  ")
    add_page_field(footer_p)
    for run in footer_p.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(90, 90, 90)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("Senatla Ops GitHub Delivery Backlog")
    remove_paragraph_borders(title)
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("Repository audit, project board controls, and implementation-ready work")
    date = doc.add_paragraph()
    date.add_run("Updated 31 October 2026").bold = True
    intro = doc.add_paragraph()
    intro.add_run("Purpose. ").bold = True
    intro.add_run("This report converts the repository audit into an accountable delivery plan. Broad issues remain outcome epics; the new tasks define smaller work boundaries with acceptance criteria, dependencies, estimates, ownership, milestones, and board fields.")
    conclusion = doc.add_paragraph()
    conclusion.add_run("Current conclusion. ").bold = True
    conclusion.add_run("The immediate release risk is governance and proof, not a lack of feature ideas. The next work should align the release branches, prove hosted configuration, complete persisted operational boundaries, and then run cross-role UAT.")

    doc.add_page_break()
    add_heading(doc, "Contents", 1)
    for item in (
        "1 Current delivery state",
        "2 Delivery roadmap",
        "3 Implementation ready tasks",
        "4 Critical execution sequence",
        "5 Definition of done",
        "6 Commercial viability review",
        "7 Three pass review",
    ):
        p = doc.add_paragraph(item)
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.space_after = Pt(4)

    add_heading(doc, "1 Current delivery state", 1)
    state_table = doc.add_table(rows=1, cols=4)
    state_table.autofit = False
    headers = ["Issues organized", "New tasks", "Milestones", "Accountable owner"]
    values = ["29", "12", "5", "OhhttooeXCIV"]
    for i, text in enumerate(headers):
        cell = state_table.rows[0].cells[i]
        cell.text = text
        set_cell_shading(cell, NAVY)
        set_cell_borders(cell)
        set_cell_margin(cell)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    row = state_table.add_row()
    for i, text in enumerate(values):
        cell = row.cells[i]
        cell.text = text
        set_cell_borders(cell)
        set_cell_margin(cell, top=150, bottom=150)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.size = Pt(15 if i < 3 else 10)
            run.font.color.rgb = RGBColor(23, 54, 93)
    for text in (
        "Every repository issue is assigned to OhhttooeXCIV; the existing collaborator remains on issues 19 and 20.",
        "Issues 4 through 20 and tasks 86 through 97 are on the Senatla Ops project board.",
        "Corrected Target Date, Delivery Stage, Risk Category, and Client Decision fields are available.",
        "Pull request 82 is classified as P0 Security work in Code Review.",
    ):
        add_bullet(doc, text)

    doc.add_page_break()
    add_heading(doc, "2 Delivery roadmap", 1)
    roadmap = [
        ("M0 Governance and Access", "Branch rules, identity, data boundaries, deployment access, and client decisions", "Protected release path, access proof, negative tests, and recorded decisions"),
        ("M1 Operational Foundations", "Site-day, workforce, recovery, private evidence, and document controls", "Persisted workflows, recovery proof, audit history, and reproducible artifacts"),
        ("M2 Controlled Pilot", "Assets, inspections, maintenance, and source-backed analytics", "Clean import, safety blocks, independent approvals, and reconciled metrics"),
        ("M3 UAT and Release", "Role journeys, recovery rehearsal, production mapping, and acceptance", "Desktop and mobile proof, restore rehearsal, rollback target, and sign-off"),
        ("M4 Reuse and Scale", "Reusable architecture and delivery workflows", "Domain-neutral packages, consumer proof, and adoption guidance"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    widths = [1.55, 2.75, 2.35]
    for i, text in enumerate(("Milestone", "Purpose", "Exit evidence")):
        cell = table.rows[0].cells[i]
        cell.width = Inches(widths[i])
        cell.text = text
        set_cell_shading(cell, NAVY)
        set_cell_borders(cell)
        set_cell_margin(cell)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
    set_repeat_table_header(table.rows[0])
    for idx, values_ in enumerate(roadmap):
        row = table.add_row()
        for i, text in enumerate(values_):
            cell = row.cells[i]
            cell.width = Inches(widths[i])
            cell.text = text
            set_cell_borders(cell)
            set_cell_margin(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if idx % 2:
                set_cell_shading(cell, PALE_BLUE)

    doc.add_page_break()
    add_heading(doc, "3 Implementation ready tasks", 1)
    summary = doc.add_table(rows=1, cols=5)
    summary.autofit = False
    widths = [0.55, 0.62, 3.25, 1.0, 1.0]
    for i, text in enumerate(("Issue", "Priority", "Task", "Estimate", "Stage")):
        cell = summary.rows[0].cells[i]
        cell.width = Inches(widths[i])
        cell.text = text
        set_cell_shading(cell, NAVY)
        set_cell_borders(cell)
        set_cell_margin(cell)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(8.5)
    set_repeat_table_header(summary.rows[0])
    for idx, (number, priority, _parent, task, estimate, stage, _criteria) in enumerate(TASKS):
        row = summary.add_row()
        values_ = [str(number), priority, task, estimate, stage.split()[0]]
        for i, text in enumerate(values_):
            cell = row.cells[i]
            cell.width = Inches(widths[i])
            cell.text = text
            set_cell_borders(cell)
            set_cell_margin(cell, top=80, bottom=80)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if idx % 2:
                set_cell_shading(cell, LIGHT_GRAY)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if i in (0, 1, 3, 4) else WD_ALIGN_PARAGRAPH.LEFT
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(8.5)

    for number, priority, parent, task, estimate, stage, criteria in TASKS:
        add_heading(doc, f"Issue {number} {task}", 2)
        p = doc.add_paragraph()
        p.add_run(f"Priority {priority}  |  Parent epic {parent}  |  {stage}  |  Estimate {estimate}").bold = True
        p.paragraph_format.keep_with_next = True
        for item in criteria:
            add_bullet(doc, item, keep_with_next=True)
        link = doc.add_paragraph()
        link.add_run("GitHub issue: ").bold = True
        link.add_run(f"https://github.com/Sankofa-Digital-Studio/senatla-ops-internal-automation-solution/issues/{number}")

    add_heading(doc, "4 Critical execution sequence", 1)
    sequence = [
        "Complete issue 87 so new work targets dev and main remains a reviewed production branch.",
        "Resolve issue 88 to prove hosted runtime configuration and controlled QA access.",
        "Complete issues 86, 89, 94, and 95 as the operational foundation.",
        "Complete issues 90, 91, 92, and 93 for the controlled pilot.",
        "Run issue 96 after the dependent role workflows are stable.",
        "Complete issue 97 after the Senatla-specific gate is stable enough to extract without weakening it.",
    ]
    for item in sequence:
        p = doc.add_paragraph(item, style="List Number")
        p.paragraph_format.space_after = Pt(4)

    add_heading(doc, "5 Definition of done", 1)
    for item in (
        "A focused pull request links to the issue and demonstrates the acceptance criteria.",
        "New mutations create required audit evidence and authorization changes include negative tests.",
        "Rendered changes have desktop and 390 px proof against the reviewed commit.",
        "Schema changes have clean reset and database-contract evidence.",
        "Deployment changes record the approved commit, environment, current deployment, and rollback target.",
        "Durable documents avoid stale immutable preview URLs and the public Senatla Trading website remains untouched.",
    ):
        add_bullet(doc, item)

    add_heading(doc, "6 Commercial viability review", 1)
    doc.add_paragraph("MaintainX, Fleetio, and HCSS Equipment360 already provide mature work-order, fleet, and equipment-maintenance capabilities. Senatla Ops should not reproduce a broad CMMS or fleet suite.")
    doc.add_paragraph("Its viable boundary is a source-linked operational record for one operating company. People, safety, attendance, payroll control, site synchronization, assets, maintenance, maker-checker decisions, and audit evidence share one authorization and audit model. The differentiator is the explainable operational heartbeat: each outcome can be traced to an actor, site, time, approval, and source record.")
    sources = (
        "MaintainX work order management: https://www.getmaintainx.com/use-cases/work-order-management",
        "Fleetio asset management: https://www.fleetio.com/features/asset-management-software",
        "HCSS Equipment360: https://www.hcss.com/equipment360/what-is-equipment360/",
    )
    for source in sources:
        add_bullet(doc, source)

    doc.add_page_break()
    add_heading(doc, "7 Three pass review", 1)
    add_heading(doc, "Pass 1 Repository gaps", 2)
    doc.add_paragraph("The initial scan found broad epics, incomplete board coverage, inconsistent repository metadata, narrow cross-role browser proof, and unfinished persistence boundaries for site-day, inspection, evidence, and reporting workflows.")
    add_heading(doc, "Pass 2 Duplication challenge", 2)
    doc.add_paragraph("Authentication, negative RLS contracts, deterministic verification, attendance retry, readiness confirmation, and parts of the offline outbox already exist. New issues for those features would duplicate closed or active work.")
    add_heading(doc, "Pass 3 Reconciled backlog", 2)
    doc.add_paragraph("The final tasks target observable remaining boundaries: live GitHub configuration, hosted deployment proof, missing persisted fields, immutable history, controlled import, inspection and safety-block contracts, database-enforced return approval, private evidence, source-backed metrics, reproducible documents, cross-role UAT, and reusable CI.")

    settings = doc.settings._element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)
    return doc


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
